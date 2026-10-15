"""Small helper layer over python-pptx for an engineering-style deck.

Thin-line diagrams, ruled tables, status markers. All positions in inches.
"""
import copy
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- palette
INK = "1D2733"      # primary text
MUTED = "5B6776"    # secondary text
FAINT = "8C97A3"    # tertiary / captions
LINE = "7F8B97"     # diagram outlines
RULE = "D3D9DF"     # table rules, frames
BG2 = "F3F5F7"      # very light gray fill
BLUE = "1F5C99"     # clock / reference, primary accent
BLUE_L = "E8F0F8"
TEAL = "0E7C76"     # RF signal
TEAL_L = "E4F2F1"
AMBER = "A86A0B"    # trigger / control, open question
AMBER_L = "FBF3E6"
RED = "A23B32"      # used very sparingly (risk)
WHITE = "FFFFFF"

FONT = "Calibri"
FONT_LIGHT = "Calibri Light"

SLIDE_W, SLIDE_H = 13.333, 7.5


def rgb(h):
    return RGBColor.from_string(h)


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SLIDE_W)
        self.prs.slide_height = Inches(SLIDE_H)
        self.layout = self.prs.slide_layouts[5]  # Title Only
        self.n = 0
        self.notes = {}

    def slide(self, title, source=None, number=True):
        s = self.prs.slides.add_slide(self.layout)
        self.n += 1
        t = s.shapes.title
        t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.38), Inches(12.1), Inches(0.75)
        tf = t.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.text = title
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        for r in p.runs:
            r.font.size = Pt(30)
            r.font.name = FONT_LIGHT
            r.font.bold = False
            r.font.color.rgb = rgb(INK)
        if source:
            text(s, 0.6, 7.02, 10.6, 0.3, source, size=10, color=FAINT, anchor="b")
        if number:
            text(s, 11.9, 7.02, 0.8, 0.3, str(self.n), size=10, color=FAINT, align="r", anchor="b")
        return s

    def save(self, path):
        self.prs.save(path)


# ---------------------------------------------------------------- text
def _apply_run(r, size, color, bold, italic, font, sup=False):
    f = r.font
    f.size = Pt(size)
    f.color.rgb = rgb(color)
    f.bold = bold
    f.italic = italic
    f.name = font
    if sup:
        r._r.get_or_add_rPr().set("baseline", "30000")


def fill_tf(tf, content, size=14, color=INK, bold=False, italic=False, align="l",
            font=FONT, bullets=False, para_space=0, line_spacing=None, level_indent=0.22):
    """content: str | list[para]; para: str | list[run]; run: str | (str, dict).
    A para may also be a dict {"runs": [...], "bullet": bool, "level": int, "size":..}"""
    if isinstance(content, str):
        content = content.split("\n")
    tf.clear()
    first = True
    for para in content:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        pbul = bullets
        psize = size
        pcolor = color
        pspace = para_space
        if isinstance(para, dict):
            pbul = para.get("bullet", bullets)
            psize = para.get("size", size)
            pcolor = para.get("color", color)
            pspace = para.get("space", para_space)
            runs = para["runs"]
        else:
            runs = para
        if isinstance(runs, str):
            runs = [runs]
        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
        if pspace:
            p.space_after = Pt(pspace)
        if line_spacing:
            p.line_spacing = line_spacing
        pPr = p._p.get_or_add_pPr()
        if pbul:
            pPr.set("marL", str(int(Inches(level_indent))))
            pPr.set("indent", str(-int(Inches(level_indent))))
            bu = etree.SubElement(pPr, qn("a:buFont"))
            bu.set("typeface", "Arial")
            bc = etree.SubElement(pPr, qn("a:buChar"))
            bc.set("char", "–" if pbul == "dash" else "•")
        else:
            etree.SubElement(pPr, qn("a:buNone"))
        for run in runs:
            if isinstance(run, tuple):
                txt, o = run
            else:
                txt, o = run, {}
            r = p.add_run()
            r.text = txt
            _apply_run(r, o.get("size", psize), o.get("color", pcolor), o.get("bold", bold),
                       o.get("italic", italic), o.get("font", font), o.get("sup", False))
        end = p._p.get_or_add_endParaRPr()
        end.set("sz", str(int(psize * 100)))
    return tf


def text(slide, x, y, w, h, content, size=14, color=INK, bold=False, italic=False, align="l",
         anchor="t", font=FONT, bullets=False, para_space=0, margin=0, line_spacing=None,
         name=None, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    m = Inches(margin)
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = m
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    fill_tf(tf, content, size, color, bold, italic, align, font, bullets, para_space, line_spacing)
    return tb


# ---------------------------------------------------------------- shapes
def _line_fmt(shape_line, color, lw, dash):
    if color is None:
        shape_line.fill.background()
        return
    shape_line.color.rgb = rgb(color)
    shape_line.width = Pt(lw)
    if dash:
        ln = shape_line._get_or_add_ln()
        pd = ln.find(qn("a:prstDash"))
        if pd is None:
            pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", dash)


def rect(slide, x, y, w, h, fill=None, line=LINE, lw=0.75, dash=None, content=None, size=12,
         color=INK, bold=False, align="c", anchor="m", shape=MSO_SHAPE.RECTANGLE, margin=0.06,
         name=None, italic=False, font=FONT, bullets=False, para_space=0, adj=None):
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        sh.name = name
    if adj is not None:
        sh.adjustments[0] = adj
    sh.shadow.inherit = False
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    _line_fmt(sh.line, line, lw, dash)
    tf = sh.text_frame
    tf.word_wrap = True
    m = Inches(margin)
    tf.margin_left = tf.margin_right = m
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    if content is not None:
        fill_tf(tf, content, size, color, bold, italic, align, font, bullets, para_space)
    else:
        fill_tf(tf, "", size, color)
    return sh


def line(slide, x1, y1, x2, y2, color=LINE, lw=1.0, dash=None, head=False, tail=False, name=None):
    """head = arrow at start (x1,y1); tail = arrow at end (x2,y2)."""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    if name:
        c.name = name
    _line_fmt(c.line, color, lw, dash)
    ln = c.line._get_or_add_ln()
    for flag, tag in ((head, "a:headEnd"), (tail, "a:tailEnd")):
        if flag:
            e = etree.SubElement(ln, qn(tag))
            e.set("type", "triangle")
            e.set("w", "sm")
            e.set("len", "sm")
    return c


def path(slide, pts, color=LINE, lw=1.0, dash=None, start_arrow=False, end_arrow=False):
    """Orthogonal polyline as consecutive straight connectors."""
    segs = []
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        segs.append(line(slide, x1, y1, x2, y2, color, lw, dash,
                         head=(start_arrow and i == 0), tail=(end_arrow and i == len(pts) - 2)))
    return segs


def dot(slide, x, y, d=0.09, fill=INK, line_color=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - d / 2), Inches(y - d / 2), Inches(d), Inches(d))
    sh.shadow.inherit = False
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    if line_color:
        _line_fmt(sh.line, line_color, 0.75, None)
    else:
        sh.line.fill.background()
    return sh


# status markers: confirmed = filled teal disc, inferred = blue ring,
# open = amber dashed ring, speculative = gray dotted ring
STATUS = {
    "C": ("Confirmed", TEAL),
    "I": ("Engineering inference", BLUE),
    "O": ("Open — needs bench test", AMBER),
    "S": ("Speculative", FAINT),
}


def status_mark(slide, x, y, kind, d=0.15):
    """(x, y) = centre."""
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - d / 2), Inches(y - d / 2), Inches(d), Inches(d))
    sh.shadow.inherit = False
    name, col = STATUS[kind]
    sh.name = f"status {name}"
    if kind == "C":
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(col)
        sh.line.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(WHITE)
        _line_fmt(sh.line, col, 1.75 if kind == "I" else 1.25,
                  {"I": None, "O": "dash", "S": "sysDot"}[kind])
    return sh


def status_legend(slide, x, y, kinds=("C", "I", "O"), size=10, gap=0.25):
    cx = x
    for k in kinds:
        status_mark(slide, cx + 0.08, y + 0.11, k, d=0.13)
        label = STATUS[k][0]
        w = 0.062 * len(label) * size / 10 + 0.05
        text(slide, cx + 0.22, y, w, 0.22, label, size=size, color=MUTED, anchor="m")
        cx += 0.22 + w + gap
    return cx


# ---------------------------------------------------------------- tables
NO_STYLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _cell_borders(cell, top=None, bottom=None, left=None, right=None, fill=None):
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB", "a:solidFill", "a:noFill"):
        for e in tcPr.findall(qn(tag)):
            tcPr.remove(e)

    def mk(tag, spec):
        ln = etree.SubElement(tcPr, qn(tag))
        if spec is None:
            ln.set("w", "0")
            etree.SubElement(ln, qn("a:noFill"))
        else:
            col, w = spec
            ln.set("w", str(int(Pt(w))))
            sf = etree.SubElement(ln, qn("a:solidFill"))
            c = etree.SubElement(sf, qn("a:srgbClr"))
            c.set("val", col)
        return ln

    mk("a:lnL", left)
    mk("a:lnR", right)
    mk("a:lnT", top)
    mk("a:lnB", bottom)
    if fill:
        sf = etree.SubElement(tcPr, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr"))
        c.set("val", fill)
    else:
        etree.SubElement(tcPr, qn("a:noFill"))


def table(slide, x, y, col_w, rows, row_h=0.32, size=12, header=True, header_fill=None,
          header_color=MUTED, header_size=None, col_align=None, first_col_color=INK,
          first_col_bold=False, zebra=None, valign="m", name=None, row_heights=None,
          cell_margin=0.07, rule=RULE, header_rule=LINE, color=INK):
    """rows: list of lists; a cell may be str or list-of-paras (fill_tf content) or
    dict {"c": content, "color":..., "bold":..., "fill":..., "align":...}."""
    nr, nc = len(rows), len(col_w)
    shp = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(sum(col_w)), Inches(row_h * nr))
    if name:
        shp.name = name
    tbl = shp.table
    tblPr = tbl._tbl.tblPr
    for a in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(a, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE
    for j, w in enumerate(col_w):
        tbl.columns[j].width = Inches(w)
    for i in range(nr):
        tbl.rows[i].height = Inches(row_heights[i] if row_heights else row_h)
    for i, row in enumerate(rows):
        for j in range(nc):
            val = row[j] if j < len(row) else ""
            cell = tbl.cell(i, j)
            opts = {}
            if isinstance(val, dict):
                opts = val
                val = val.get("c", "")
            is_head = header and i == 0
            cfill = opts.get("fill")
            if is_head and header_fill:
                cfill = header_fill
            elif zebra and not is_head and (i % 2 == 0):
                cfill = cfill or zebra
            top = None
            bottom = (header_rule, 1.0) if is_head else (rule, 0.75)
            if i == nr - 1 and not is_head:
                bottom = (rule, 0.75)
            _cell_borders(cell, top=top, bottom=bottom, fill=cfill)
            cell.margin_left = cell.margin_right = Inches(cell_margin)
            cell.margin_top = Inches(0.035)
            cell.margin_bottom = Inches(0.035)
            cell.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[
                opts.get("valign", valign)]
            al = opts.get("align") or (col_align[j] if col_align else "l")
            if is_head:
                fill_tf(cell.text_frame, val, header_size or size, opts.get("color", header_color),
                        opts.get("bold", False), False, al)
            else:
                fill_tf(cell.text_frame, val, opts.get("size", size),
                        opts.get("color", first_col_color if j == 0 else color),
                        opts.get("bold", first_col_bold and j == 0), opts.get("italic", False), al)
    return shp


# ---------------------------------------------------------------- misc
def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def picture(slide, path_, x, y, w=None, h=None, name=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    p = slide.shapes.add_picture(path_, Inches(x), Inches(y), **kw)
    if name:
        p.name = name
    return p
